#!/usr/bin/env python
# =============================================================================
#  Systematic review pipeline: uncertainty-aware and trustworthy LLMs in education
#  (IEEE Transactions on Learning Technologies survey)
#
#  Stages (env STAGES, default "harvest,merge,llm,sheet"):
#    harvest : search OpenAlex, arXiv and ERIC through their public APIs (Internet ON)
#    merge   : read your database exports (Scopus, Web of Science, IEEE Xplore,
#              ACM DL, ERIC) from /kaggle/input, add the API records, remove
#              duplicates, compute keyword-rule flags
#    llm     : optional AI pre-screen (Qwen2.5-7B-Instruct, 4-bit; one forward
#              pass per record; probability that the record meets the criteria)
#    sheet   : Excel screening workbook + PRISMA counts + search log
#    kappa   : run later on the filled workbook (SHEET=/kaggle/input/.../screening.xlsx)
#              -> Cohen's kappa, conflicts, and agreement of the AI pre-screen
#
#  Env: STAGES, OPENALEX_KEY (optional, free key), EMAIL (optional), FAST=1,
#       LLM_MODEL, SHEET, OUT_DIR, DATA_DIR
#  Output: /kaggle/working/survey_screening_<date>.zip
# =============================================================================
import os, re, sys, json, time, glob, math, random, zipfile, shutil, traceback, unicodedata
from datetime import datetime
import numpy as np
import pandas as pd

T0 = time.time()
ON_KAGGLE = os.path.isdir("/kaggle/working")
OUT = os.environ.get("OUT_DIR", "/kaggle/working/survey" if ON_KAGGLE else os.path.abspath("./survey_out"))
DATA_DIR = os.environ.get("DATA_DIR", "/kaggle/input" if ON_KAGGLE else os.path.abspath("./exports"))
os.makedirs(OUT, exist_ok=True)
FAST = os.environ.get("FAST", "0") not in ("0", "", "false")
STAGES = [s.strip() for s in os.environ.get("STAGES", "harvest,merge,llm,sheet").split(",") if s.strip()]
TODAY = datetime.now().strftime("%Y-%m-%d")
YEAR_FROM, YEAR_TO = 2020, datetime.now().year
SEED = 2026


def log(*a):
    line = "[%6.1f min] %s" % ((time.time() - T0) / 60, " ".join(str(x) for x in a))
    print(line, flush=True)
    with open(os.path.join(OUT, "run.log"), "a") as f:
        f.write(line + "\n")


# ----------------------------------------------------------------------------- search terms
# Explicit word forms (no wildcards), so the same terms work in every API.
A = ["large language model", "large language models", "LLM", "LLMs", "generative AI",
     "generative artificial intelligence", "GPT", "GPT-3", "GPT-4", "ChatGPT", "pretrained language model",
     "pretrained language models", "pre-trained language model", "pre-trained language models"]
# Version 2 (after the pilot run): core uncertainty terms only. Bare "confidence", "uncertain",
# "trustworthy", "hallucination" and "over-reliance" matched thousands of papers that only mention these words.
B = ["uncertainty", "uncertainties", "confidence score", "confidence scores", "confidence estimation", "model confidence",
     "confidence calibration", "verbalized confidence", "verbalised confidence", "overconfidence", "overconfident",
     "calibration", "calibrated", "miscalibration", "miscalibrated", "conformal", "abstain", "abstention",
     "selective prediction", "selective classification", "deferral", "learning to defer", "human-in-the-loop",
     "teacher-in-the-loop", "hallucination detection", "hallucination rate", "hallucination rates",
     "hallucination mitigation", "detect hallucinations", "detecting hallucinations", "factual consistency",
     "trust calibration", "appropriate reliance"]
C = ["education", "educational", "student", "students", "learner", "learners", "teacher", "teachers", "tutor",
     "tutoring", "automated scoring", "grading", "automated feedback", "formative assessment", "summative assessment",
     "automated assessment", "student assessment", "educational assessment", "classroom", "MOOC", "learning analytics"]
NOT = ["knowledge distillation", "teacher-student"]      # "teacher"/"student" in model distillation, not education


def q_or(terms):
    return "(" + " OR ".join('"%s"' % t if (" " in t or "-" in t) else t for t in terms) + ")"


def norm_text(s):
    s = unicodedata.normalize("NFKD", str(s or "")).lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def term_hits(text, terms):
    t = " " + norm_text(text) + " "
    return sorted({x for x in terms if " " + norm_text(x) + " " in t})


# ----------------------------------------------------------------------------- HTTP helper
def http_get(url, params=None, headers=None, tries=5, timeout=60):
    import requests
    for k in range(tries):
        try:
            r = requests.get(url, params=params, headers=headers, timeout=timeout)
            if r.status_code == 200:
                return r
            if r.status_code in (429, 500, 502, 503, 504):
                if k == tries - 1:
                    log("   HTTP %d after %d tries: %s" % (r.status_code, tries, r.text[:300]))
                    if r.status_code == 429 and "openalex" in url:
                        log("!! OpenAlex daily quota or rate limit reached. Without a key the free budget is small and "
                            "is shared by all runs on the same day. Add a free API key as the Kaggle Secret OPENALEX_KEY "
                            "(the notebook must print 'OpenAlex key loaded') or wait until the quota resets at 00:00 UTC.")
                    return r
                time.sleep(3 * (k + 1))
                continue
            log("   HTTP %d: %s" % (r.status_code, r.text[:200]))
            return r
        except Exception as e:
            log("   request error (%s), retrying" % str(e)[:120])
            time.sleep(3 * (k + 1))
    return None


# ----------------------------------------------------------------------------- harvest: OpenAlex
def openalex_abstract(inv):
    if not inv:
        return ""
    pos = []
    for w, idx in inv.items():
        for i in idx:
            pos.append((i, w))
    return " ".join(w for _, w in sorted(pos))


def harvest_openalex(search_log):
    key = os.environ.get("OPENALEX_KEY", "").strip()
    email = os.environ.get("EMAIL", "").strip()
    query = "%s AND %s AND %s NOT %s" % (q_or(A), q_or(B), q_or(C), q_or(NOT))
    base = "https://api.openalex.org/works"
    sel = ("id,doi,title,display_name,publication_year,type,primary_location,authorships,"
           "abstract_inverted_index,cited_by_count,language")
    headers = {"Authorization": "Bearer " + key} if key else {}
    log("OpenAlex: %s" % ("using your API key" if key else "no API key (small shared daily quota)"))
    modes = [("filter", {"filter": "title_and_abstract.search:%s,publication_year:%d-%d" % (query, YEAR_FROM, YEAR_TO)}),
             ("search.title_abstract_keywords", {"search.title_abstract_keywords": query,
                                                 "filter": "publication_year:%d-%d" % (YEAR_FROM, YEAR_TO)})]
    rows = []
    def get_json(params):
        """GET with retries; also retries when a response arrives truncated (invalid JSON)."""
        for k in range(4):
            r = http_get(base, params, headers)
            if r is None:
                return None, None
            if r.status_code != 200:
                return r.status_code, None
            try:
                return 200, r.json()
            except Exception:
                log("   truncated OpenAlex response; retrying (%d)" % (k + 1))
                time.sleep(3 * (k + 1))
        return 200, None
    def add(js):
        for w in js.get("results", []):
            loc = (w.get("primary_location") or {})
            src = (loc.get("source") or {})
            rows.append(dict(title=w.get("display_name") or w.get("title") or "",
                             abstract=openalex_abstract(w.get("abstract_inverted_index")),
                             doi=(w.get("doi") or ""), year=w.get("publication_year"),
                             venue=src.get("display_name") or "", type=w.get("type") or "",
                             authors="; ".join((a.get("author") or {}).get("display_name", "")
                                               for a in (w.get("authorships") or [])[:12]),
                             citations=w.get("cited_by_count"), language=w.get("language") or "",
                             url=w.get("id") or "", source="OpenAlex"))
    limit = 300 if FAST else 20000
    for mname, mparams in modes:
        # cursor paging needs a key; without one, numbered pages work up to 10,000 records
        params = dict(mparams, select=sel, per_page=50)
        params.update({"cursor": "*"} if key else {"page": 1})
        if email:
            params["mailto"] = email
        code, js = get_json(params)
        if js is None:
            log("   OpenAlex mode '%s' not accepted (%s); trying the next one" % (mname, code))
            continue
        total = js.get("meta", {}).get("count", 0)
        log("OpenAlex (%s): %d matching works" % (mname, total))
        search_log.append(dict(source="OpenAlex (API)", date=TODAY, query=query, field="title and abstract",
                               years="%d-%d" % (YEAR_FROM, YEAR_TO), hits=total))
        if not key and total > 10000:
            log("!! %d matches but only the first 10,000 can be paged without a free OpenAlex API key" % total)
        while js is not None:
            add(js)
            if len(rows) >= min(limit, total) or not js.get("results"):
                break
            if key:
                nxt = js.get("meta", {}).get("next_cursor")
                if not nxt:
                    break
                params["cursor"] = nxt
            else:
                params["page"] += 1
                if params["page"] > 200:
                    break
            time.sleep(0.12)
            code, js = get_json(params)
            for wait in (30, 60, 120):          # network hiccup: pause and resume from the same page
                if js is not None:
                    break
                log("   OpenAlex not responding at %d records; waiting %d s and resuming" % (len(rows), wait))
                time.sleep(wait)
                code, js = get_json(params)
            if js is None:
                log("!! OpenAlex stopped at %d records (status %s); the records so far are kept" % (len(rows), code))
        break
    log("OpenAlex: %d records downloaded" % len(rows))
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------- harvest: arXiv
def harvest_arxiv(search_log):
    import xml.etree.ElementTree as ET
    def blk(terms):
        return "(" + " OR ".join('abs:"%s"' % t if (" " in t or "-" in t) else "abs:%s" % t for t in terms) + ")"
    cats = "(cat:cs.CL OR cat:cs.AI OR cat:cs.CY OR cat:cs.HC OR cat:cs.LG OR cat:cs.IR OR cat:cs.SE)"
    q = "%s AND %s AND %s AND %s AND submittedDate:[%d01010000 TO %d12312359] ANDNOT %s" % (
        blk(A), blk(B), blk(C), cats, YEAR_FROM, YEAR_TO, blk(NOT))
    ns = {"a": "http://www.w3.org/2005/Atom", "o": "http://a9.com/-/spec/opensearch/1.1/"}
    rows, start, step = [], 0, 200
    limit = 200 if FAST else 10000
    total = None
    while True:
        r = http_get("https://export.arxiv.org/api/query",
                     {"search_query": q, "start": start, "max_results": step, "sortBy": "submittedDate",
                      "sortOrder": "descending"})
        if r is None or r.status_code != 200:
            break
        root = ET.fromstring(r.content)
        if total is None:
            el = root.find("o:totalResults", ns)
            total = int(el.text) if el is not None else 0
            log("arXiv: %d matching e-prints" % total)
            search_log.append(dict(source="arXiv (API)", date=TODAY, query=q, field="abstract",
                                   years="%d-%d" % (YEAR_FROM, YEAR_TO), hits=total))
        ents = root.findall("a:entry", ns)
        if not ents:
            break
        for e in ents:
            g = lambda tag: (e.find("a:" + tag, ns).text or "").strip() if e.find("a:" + tag, ns) is not None else ""
            doi_el = e.find("{http://arxiv.org/schemas/atom}doi")
            pub = g("published")
            rows.append(dict(title=re.sub(r"\s+", " ", g("title")), abstract=re.sub(r"\s+", " ", g("summary")),
                             doi=(doi_el.text if doi_el is not None else ""), year=int(pub[:4]) if pub else None,
                             venue="arXiv", type="preprint",
                             authors="; ".join((a.find("a:name", ns).text or "") for a in e.findall("a:author", ns)[:12]),
                             citations=None, language="en", url=g("id"), source="arXiv"))
        start += step
        if start >= min(total or 0, limit):
            break
        time.sleep(3.1)            # arXiv asks for at most one request every 3 seconds
    log("arXiv: %d records downloaded" % len(rows))
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------- harvest: ERIC
def harvest_eric(search_log):
    """ERIC (education database of the US Institute of Education Sciences); free public API, no key."""
    def blk(terms):
        return "(" + " OR ".join('"%s"' % t if (" " in t or "-" in t) else t for t in terms) + ")"
    q = "%s AND %s NOT %s" % (blk(A), blk(B), blk(NOT))   # every ERIC record is educational, so block C is not needed
    rows, start, step = [], 0, 200
    limit = 200 if FAST else 5000
    total = None
    while True:
        r = http_get("https://api.ies.ed.gov/eric/", {"search": q, "format": "json", "rows": step, "start": start,
                     "fields": "id,title,author,description,publicationdateyear,source,publicationtype,url,peerreviewed"})
        if r is None or r.status_code != 200:
            break
        resp = r.json().get("response", {})
        if total is None:
            total = int(resp.get("numFound", 0))
            log("ERIC: %d matching records (all years; filtered to %d-%d later)" % (total, YEAR_FROM, YEAR_TO))
            search_log.append(dict(source="ERIC (API)", date=TODAY, query=q, field="all fields", years="all", hits=total))
        docs = resp.get("docs", [])
        if not docs:
            break
        for d in docs:
            j = lambda v: "; ".join(v) if isinstance(v, list) else (v or "")
            rows.append(dict(title=j(d.get("title")), abstract=j(d.get("description")), doi="",
                             year=d.get("publicationdateyear"), venue=j(d.get("source")),
                             type=j(d.get("publicationtype")), authors=j(d.get("author"))[:400], citations=None,
                             language="en", url="https://eric.ed.gov/?id=%s" % d.get("id", ""), source="ERIC"))
        start += step
        if start >= min(total, limit):
            break
        time.sleep(0.5)
    log("ERIC: %d records downloaded" % len(rows))
    return pd.DataFrame(rows)


def stage_harvest():
    slog = []
    frames = []
    for name, fn in (("OpenAlex", harvest_openalex), ("arXiv", harvest_arxiv), ("ERIC", harvest_eric)):
        try:
            frames.append(fn(slog))
        except Exception as e:
            log("!! %s harvest failed: %s" % (name, e))
            traceback.print_exc()
    frames = [f for f in frames if f is not None and len(f)]
    api = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=["title", "source"])
    if not len(api):
        log("!! no records from the APIs. Is Internet switched ON in the notebook settings?")
    api.to_csv(os.path.join(OUT, "api_records.csv"), index=False)
    pd.DataFrame(slog).to_csv(os.path.join(OUT, "search_log_api.csv"), index=False)
    log("API records saved: %d" % len(api))


# ----------------------------------------------------------------------------- import database exports
COLMAP = {
    "title": ["Title", "Document Title", "Article Title", "TI", "title", "T1"],
    "abstract": ["Abstract", "AB", "abstract", "Description", "N2"],
    "doi": ["DOI", "DI", "doi"],
    "year": ["Year", "Publication Year", "PY", "year", "Publication Date", "Publication year"],
    "venue": ["Source title", "Source Title", "Publication Title", "SO", "Journal", "journal", "Publication Name"],
    "authors": ["Authors", "Author Full Names", "AU", "Author", "author", "Author(s)"],
    "type": ["Document Type", "DT", "Document Identifier", "Publication Type", "Type"],
    "keywords": ["Author Keywords", "DE", "Author keywords", "keywords", "Keywords"],
}


def guess_source(path, cols=()):
    n = os.path.basename(path).lower()
    for k, v in (("scopus", "Scopus"), ("wos", "Web of Science"), ("savedrecs", "Web of Science"),
                 ("webofscience", "Web of Science"), ("ieee", "IEEE Xplore"), ("acm", "ACM DL"), ("eric", "ERIC")):
        if k in n:
            return v
    cols = set(cols)
    if "EID" in cols or "Source title" in cols:
        return "Scopus"
    if "UT" in cols or ("TI" in cols and "SO" in cols):
        return "Web of Science"
    if "Document Title" in cols and ("IEEE Terms" in cols or "PDF Link" in cols):
        return "IEEE Xplore"
    return "Other export"


def from_table(df, path):
    out = pd.DataFrame()
    for k, cands in COLMAP.items():
        c = next((c for c in cands if c in df.columns), None)
        out[k] = df[c].astype(str).replace({"nan": ""}) if c else ""
    out["source"] = guess_source(path, df.columns)
    return out


def parse_bib(text):
    recs, i = [], 0
    for m in re.finditer(r"@(\w+)\s*\{", text):
        start = m.end()
        depth, j = 1, start
        while j < len(text) and depth:
            depth += {"{": 1, "}": -1}.get(text[j], 0)
            j += 1
        body = text[start:j - 1]
        fields = {}
        k = body.find(",")
        body = body[k + 1:]
        for fm in re.finditer(r"(\w+)\s*=\s*", body):
            pos = fm.end()
            if pos >= len(body):
                continue
            if body[pos] == "{":
                d, q = 1, pos + 1
                while q < len(body) and d:
                    d += {"{": 1, "}": -1}.get(body[q], 0)
                    q += 1
                val = body[pos + 1:q - 1]
            elif body[pos] == '"':
                q = body.find('"', pos + 1)
                val = body[pos + 1:q]
            else:
                q = re.search(r"[,\n]", body[pos:])
                val = body[pos:pos + (q.start() if q else len(body))]
            fields.setdefault(fm.group(1).lower(), re.sub(r"[{}]", "", val).strip())
        recs.append(dict(title=fields.get("title", ""), abstract=fields.get("abstract", ""), doi=fields.get("doi", ""),
                         year=fields.get("year", ""), venue=fields.get("journal", fields.get("booktitle", "")),
                         authors=fields.get("author", ""), type=m.group(1).lower(),
                         keywords=fields.get("keywords", "")))
    return pd.DataFrame(recs)


def parse_ris(text):
    recs, cur = [], {}
    for line in text.splitlines():
        m = re.match(r"^([A-Z][A-Z0-9])  - ?(.*)$", line)
        if not m:
            continue
        tag, val = m.group(1), m.group(2).strip()
        if tag == "ER":
            recs.append(cur); cur = {}; continue
        cur.setdefault(tag, []).append(val)
    if cur:
        recs.append(cur)
    g = lambda r, *tags: next((" ".join(r[t]) for t in tags if t in r), "")
    return pd.DataFrame([dict(title=g(r, "TI", "T1"), abstract=g(r, "AB", "N2"), doi=g(r, "DO"),
                              year=g(r, "PY", "Y1", "DA")[:4], venue=g(r, "JO", "T2", "JF", "SO"),
                              authors="; ".join(r.get("AU", []) + r.get("A1", [])), type=g(r, "TY"),
                              keywords="; ".join(r.get("KW", []))) for r in recs])


def read_exports():
    frames, counts = [], {}
    files = [p for p in glob.glob(os.path.join(DATA_DIR, "**", "*"), recursive=True) if os.path.isfile(p)]
    files = [p for p in files if re.search(r"\.(csv|txt|tsv|xls|xlsx|bib|ris)$", p, re.I)
             and "screening" not in os.path.basename(p).lower()]
    for p in files:
        ext = p.lower().rsplit(".", 1)[-1]
        try:
            if ext == "bib":
                df = parse_bib(open(p, encoding="utf-8", errors="ignore").read())
                df["source"] = guess_source(p) if guess_source(p) != "Other export" else "ACM DL"
            elif ext == "ris":
                df = parse_ris(open(p, encoding="utf-8", errors="ignore").read())
                df["source"] = guess_source(p) if guess_source(p) != "Other export" else "ERIC"
            elif ext in ("xls", "xlsx"):
                df = from_table(pd.read_excel(p), p)
            elif ext in ("txt", "tsv"):
                raw = None
                for enc in ("utf-8-sig", "utf-16", "latin-1"):
                    try:
                        raw = pd.read_csv(p, sep="\t", dtype=str, quoting=3, index_col=False, encoding=enc,
                                          on_bad_lines="skip")
                        if any(c in raw.columns for c in COLMAP["title"]):
                            break
                    except Exception:
                        raw = None
                if raw is None:
                    continue
                df = from_table(raw, p)
            else:
                raw = pd.read_csv(p, dtype=str, encoding="utf-8-sig", on_bad_lines="skip")
                if not any(c in raw.columns for c in COLMAP["title"]):
                    continue
                df = from_table(raw, p)
        except Exception as e:
            log("!! could not read %s: %s" % (p, e))
            continue
        df = df[df.title.astype(str).str.strip().str.len() > 0]
        if not len(df):
            continue
        df["file"] = os.path.basename(p)
        frames.append(df)
        counts[df.source.iloc[0]] = counts.get(df.source.iloc[0], 0) + len(df)
        log("export %s -> %s: %d records" % (os.path.basename(p), df.source.iloc[0], len(df)))
    return (pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()), counts


# ----------------------------------------------------------------------------- deduplication
def norm_doi(d):
    d = str(d or "").strip().lower()
    d = re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", d)
    return d if d.startswith("10.") else ""


def dedupe(df):
    """Union-find over: same DOI; same normalised title; near-identical titles (ratio >= 0.95, year +-1)."""
    n = len(df)
    parent = list(range(n))
    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    def union(i, j):
        a, b = find(i), find(j)
        if a != b:
            parent[max(a, b)] = min(a, b)
    df = df.reset_index(drop=True)
    df["doi_n"] = df.doi.map(norm_doi)
    df["title_n"] = df.title.map(norm_text)
    for key in ("doi_n", "title_n"):
        first = {}
        for i, v in enumerate(df[key]):
            if not v or (key == "title_n" and len(v) < 15):
                continue
            if v in first:
                union(first[v], i)
            else:
                first[v] = i
    try:
        from rapidfuzz import fuzz
        ratio = lambda a, b: fuzz.ratio(a, b) / 100
    except Exception:
        from difflib import SequenceMatcher
        ratio = lambda a, b: SequenceMatcher(None, a, b).ratio()
    yr = pd.to_numeric(df.year.astype(str).str[:4], errors="coerce").fillna(0).astype(int).values
    blocks = {}
    for i, t in enumerate(df.title_n):
        w = t.split()
        if len(w) >= 3:
            blocks.setdefault(" ".join(w[:3]), []).append(i)
    for idx in blocks.values():
        if len(idx) < 2 or len(idx) > 300:
            continue
        for a in range(len(idx)):
            for b in range(a + 1, len(idx)):
                i, j = idx[a], idx[b]
                if find(i) == find(j) or (yr[i] and yr[j] and abs(yr[i] - yr[j]) > 1):
                    continue
                if ratio(df.title_n[i], df.title_n[j]) >= 0.95:
                    union(i, j)
    df["grp"] = [find(i) for i in range(n)]
    rank = {"Scopus": 0, "Web of Science": 1, "IEEE Xplore": 2, "ACM DL": 3, "ERIC": 4, "OpenAlex": 5, "arXiv": 9}
    out = []
    for g, sub in df.groupby("grp", sort=False):
        sub = sub.assign(_r=sub.source.map(lambda s: rank.get(s, 6)), _al=sub.abstract.astype(str).str.len())
        best = sub.sort_values(["_r", "_al"], ascending=[True, False]).iloc[0].to_dict()
        best["abstract"] = max(sub.abstract.astype(str), key=len)
        best["doi"] = next((d for d in sub.doi_n if d), "")
        best["sources"] = "; ".join(sorted(set(sub.source)))
        best["n_sources"] = sub.source.nunique()
        best["also_preprint"] = bool(("arXiv" in set(sub.source)) and len(set(sub.source)) > 1)
        best["is_preprint"] = set(sub.source) == {"arXiv"}
        out.append(best)
    res = pd.DataFrame(out).drop(columns=["_r", "_al", "grp"], errors="ignore")
    return res


def stage_merge():
    exp, counts = read_exports()
    api_p = os.path.join(OUT, "api_records.csv")
    api = pd.read_csv(api_p, dtype=str) if os.path.exists(api_p) and os.path.getsize(api_p) > 5 else pd.DataFrame()
    for s in ("OpenAlex", "arXiv", "ERIC"):
        if len(api) and (api.source == s).any():
            counts[s + " (API)"] = int((api.source == s).sum())
    allr = pd.concat([x for x in (exp, api) if len(x)], ignore_index=True)
    if not len(allr):
        raise SystemExit("No records found. Add your database exports as a Kaggle dataset, or run the harvest stage.")
    for c in ("title", "abstract", "doi", "year", "venue", "authors", "type", "keywords", "url", "citations", "language"):
        if c not in allr.columns:
            allr[c] = ""
    allr = allr.fillna("")
    n_identified = len(allr)
    yr = pd.to_numeric(allr.year.astype(str).str[:4], errors="coerce")
    out_of_range = int(((yr < YEAR_FROM) | (yr > YEAR_TO)).sum())
    allr = allr[~((yr < YEAR_FROM) | (yr > YEAR_TO))]
    ded = dedupe(allr)
    log("identified %d | outside %d-%d: %d | after de-duplication: %d (duplicates removed: %d)" %
        (n_identified, YEAR_FROM, YEAR_TO, out_of_range, len(ded), n_identified - out_of_range - len(ded)))
    text = (ded.title.astype(str) + ". " + ded.abstract.astype(str))
    ded["hits_llm"] = text.map(lambda s: ", ".join(term_hits(s, A)))
    ded["hits_uncertainty"] = text.map(lambda s: ", ".join(term_hits(s, B)))
    ded["hits_education"] = text.map(lambda s: ", ".join(term_hits(s, C)))
    ded["rule_all_three"] = (ded.hits_llm != "") & (ded.hits_uncertainty != "") & (ded.hits_education != "")
    ded["no_abstract"] = ded.abstract.astype(str).str.len() < 50
    has_not = text.map(lambda s: bool(term_hits(s, NOT)))
    # Automated eligibility check (reported in PRISMA as "marked ineligible by automation tools"): some sources
    # search all fields (ERIC) or match loosely (arXiv), so every record must contain all three term blocks in
    # its own title/abstract. Records without an abstract are kept for manual screening.
    # In model-distillation papers "teacher"/"student" name models, so such records need another education term.
    core_c = [c for c in C if c not in ("teacher", "teachers", "student", "students")]
    distill = text.map(lambda s: "distill" in s.lower()) & ~text.map(lambda s: bool(term_hits(s, core_c)))
    passed = (ded.rule_all_three & ~has_not & ~distill) | ded.no_abstract
    ded[~passed].to_csv(os.path.join(OUT, "records_excluded_automation.csv"), index=False)
    n_auto = int((~passed).sum())
    log("automated term check: %d records removed (not all three term blocks in title/abstract), %d kept" % (n_auto, int(passed.sum())))
    ded = ded[passed].copy()
    rng = np.random.default_rng(SEED)
    ded = ded.iloc[rng.permutation(len(ded))].reset_index(drop=True)      # random order: avoids order bias
    ded.insert(0, "rec_id", ["R%05d" % (i + 1) for i in range(len(ded))])
    ded.to_csv(os.path.join(OUT, "records_deduplicated.csv"), index=False)
    prisma = dict(date=TODAY, identified_by_source=counts, identified_total=n_identified,
                  removed_out_of_years=out_of_range, removed_duplicates=n_identified - out_of_range - len(ded) - n_auto,
                  removed_by_automated_term_check=n_auto, search_terms_version=2,
                  records_screened=len(ded), records_without_abstract=int(ded.no_abstract.sum()),
                  records_matching_all_three_term_blocks=int(ded.rule_all_three.sum()))
    json.dump(prisma, open(os.path.join(OUT, "prisma_counts.json"), "w"), indent=1)
    log("PRISMA counts:", json.dumps(prisma))


# ----------------------------------------------------------------------------- AI pre-screen (optional)
CRITERIA = (
    "You help screen papers for a systematic review on how large language models (LLMs) in EDUCATION handle "
    "uncertainty and reliability. Answer Yes if the paper COULD be relevant, that is, if it (a) involves an LLM, "
    "ChatGPT, GPT or another pretrained language model, (b) in an educational or learning context (students, teachers, "
    "tutoring, grading, assessment, feedback, courses, educational content), and (c) deals in any way with the "
    "model's uncertainty, confidence, calibration, reliability, errors or hallucinations, abstention, fairness of its "
    "judgments, or human review or oversight of its outputs. When in doubt, answer Yes. Answer No only if the paper "
    "is clearly outside this scope (for example not about education, not about language models, or only about "
    "people's attitudes without any concern for the model's reliability).")


def stage_llm():
    try:
        import torch
    except Exception:
        log("torch not available; skipping AI pre-screen"); return
    if not torch.cuda.is_available():
        log("no GPU; skipping AI pre-screen (use Accelerator = GPU T4 x2)"); return
    p = os.path.join(OUT, "records_deduplicated.csv")
    df = pd.read_csv(p, dtype=str).fillna("")
    from transformers import AutoTokenizer, AutoModelForCausalLM
    name = os.environ.get("LLM_MODEL", "Qwen/Qwen2.5-7B-Instruct")
    four_bit = "7B" in name or "14B" in name
    kw = dict(dtype=torch.float16)
    if four_bit:
        try:
            import bitsandbytes  # noqa
        except Exception:
            os.system("pip install -q bitsandbytes")
        from transformers import BitsAndBytesConfig
        kw.update(quantization_config=BitsAndBytesConfig(load_in_4bit=True, bnb_4bit_quant_type="nf4",
                                                         bnb_4bit_compute_dtype=torch.float16), device_map={"": 0})
    t0 = time.time()
    tok = AutoTokenizer.from_pretrained(name)
    try:
        model = AutoModelForCausalLM.from_pretrained(name, **kw)
    except TypeError:
        kw["torch_dtype"] = kw.pop("dtype"); model = AutoModelForCausalLM.from_pretrained(name, **kw)
    if not four_bit:
        model = model.cuda()
    model.eval()
    log("AI pre-screen model %s loaded in %.0f s" % (name, time.time() - t0))
    yes = tok.encode("Yes", add_special_tokens=False)[0]
    no = tok.encode("No", add_special_tokens=False)[0]
    W = model.get_output_embeddings().weight
    w = (W[yes].float() - W[no].float()).detach()
    def prompt(r):
        msg = [{"role": "system", "content": CRITERIA},
               {"role": "user", "content": "Title: %s\nAbstract: %s\n\nCould this paper be relevant? Answer Yes or No."
                % (r["title"][:400], (r["abstract"] or "(no abstract)")[:2200])}]
        return tok.apply_chat_template(msg, tokenize=False, add_generation_prompt=True)
    texts = [prompt(r) for r in df.to_dict("records")]
    if FAST:                                   # quick test: score only the first 150 records
        texts = texts[:150]
    order = np.argsort([len(t) for t in texts])
    scores = np.zeros(len(texts), np.float32)
    bs = 8
    tok.padding_side = "right"
    if tok.pad_token is None:
        tok.pad_token = tok.eos_token
    with torch.no_grad():
        for i in range(0, len(texts), bs):
            idx = order[i:i + bs]
            b = tok([texts[j] for j in idx], return_tensors="pt", padding=True, truncation=True, max_length=1024).to("cuda")
            h = model.model(input_ids=b["input_ids"], attention_mask=b["attention_mask"]).last_hidden_state
            last = b["attention_mask"].sum(1) - 1
            s = (h[torch.arange(len(idx), device="cuda"), last].float() @ w.to(h.device)).cpu().numpy()
            scores[idx] = s
            if (i // bs) % 50 == 0:
                log("   pre-screen %d/%d (%.1f rec/s)" % (min(i + bs, len(texts)), len(texts), (i + bs) / (time.time() - t0)))
    p_inc = np.full(len(df), np.nan)
    p_inc[:len(scores)] = 1 / (1 + np.exp(-scores))
    df["ai_p_include"] = np.round(p_inc, 4)
    df["ai_rank"] = df.ai_p_include.rank(ascending=False, method="first")
    df["ai_suggestion"] = np.where(df.ai_p_include.isna(), "", np.where(df.ai_p_include >= 0.5, "include",
                                   np.where(df.ai_p_include >= 0.05, "maybe", "exclude")))
    df["ai_model"] = name
    df.to_csv(p, index=False)
    log("AI pre-screen done: include %d | maybe %d | exclude %d" % tuple((df.ai_suggestion == k).sum()
                                                                         for k in ("include", "maybe", "exclude")))


# ----------------------------------------------------------------------------- Excel workbook
REASONS = ["E1 not an LLM / pretrained LM", "E2 not an educational setting", "E3 no uncertainty/reliability focus",
           "E4 opinion / no method", "E5 not English", "E6 duplicate", "E7 no full text", "E8 other (note)"]


def stage_sheet():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    df = pd.read_csv(os.path.join(OUT, "records_deduplicated.csv"), dtype=str).fillna("")
    rng = np.random.default_rng(SEED)
    k = int(math.ceil(0.2 * len(df)))
    dbl = np.zeros(len(df), bool); dbl[rng.choice(len(df), k, replace=False)] = True
    df["double_screen"] = np.where(dbl, "yes", "")
    cols = ["rec_id", "title", "abstract", "year", "venue", "authors", "doi", "url", "sources", "type", "double_screen",
            "S1_decision", "S1_reason", "S1_note", "S2_decision", "S2_reason", "S2_note", "final_decision",
            "hits_llm", "hits_uncertainty", "hits_education", "rule_all_three", "no_abstract", "is_preprint",
            "ai_p_include", "ai_rank", "ai_suggestion"]
    for c in cols:
        if c not in df.columns:
            df[c] = ""
    wb = Workbook()
    ws = wb.active
    ws.title = "screening"
    ws.append(cols)
    for r in df[cols].itertuples(index=False):
        ws.append([str(v) for v in r])
    widths = {"rec_id": 9, "title": 50, "abstract": 90, "year": 6, "venue": 24, "authors": 22, "doi": 18, "url": 14,
              "sources": 16, "type": 10, "double_screen": 8, "S1_decision": 11, "S1_reason": 22, "S1_note": 18,
              "S2_decision": 11, "S2_reason": 22, "S2_note": 18, "final_decision": 12}
    hdr = PatternFill("solid", fgColor="1D2B4F")
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hdr
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        letter = cell.column_letter
        ws.column_dimensions[letter].width = widths.get(c, 12)
        if c in ("hits_llm", "hits_uncertainty", "hits_education", "rule_all_three", "no_abstract", "is_preprint",
                 "ai_p_include", "ai_rank", "ai_suggestion"):
            ws.column_dimensions[letter].hidden = True      # hidden by default to avoid anchoring the screeners
    for row in ws.iter_rows(min_row=2, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    n = len(df) + 1
    dec = DataValidation(type="list", formula1='"include,exclude,maybe"', allow_blank=True)
    rea = DataValidation(type="list", formula1="=lists!$A$1:$A$%d" % len(REASONS), allow_blank=True)
    ws.add_data_validation(dec); ws.add_data_validation(rea)
    for c in ("S1_decision", "S2_decision", "final_decision"):
        L = ws.cell(row=1, column=cols.index(c) + 1).column_letter
        dec.add("%s2:%s%d" % (L, L, n))
    for c in ("S1_reason", "S2_reason"):
        L = ws.cell(row=1, column=cols.index(c) + 1).column_letter
        rea.add("%s2:%s%d" % (L, L, n))
    wl = wb.create_sheet("lists")
    for i, r in enumerate(REASONS, 1):
        wl.cell(row=i, column=1, value=r)
    wi = wb.create_sheet("instructions")
    for line in [
        "Screening instructions (title and abstract stage)",
        "1. Screener 1 (S1) screens every row; Screener 2 (S2) screens rows with double_screen = yes (a random 20%).",
        "2. Decision: include / exclude / maybe. 'maybe' goes to full-text review. When unsure, choose maybe.",
        "3. For exclude, choose the first reason that applies (E1-E8).",
        "4. Include only if ALL hold: LLM or pretrained language model; educational setting; the paper measures,",
        "   communicates or acts on model uncertainty, confidence, calibration, reliability, hallucination, abstention or deferral.",
        "5. Do not look at the hidden AI columns until you have finished; they are used only to report AI-screener agreement.",
        "6. Disagreements are resolved by discussion; record the result in final_decision.",
        "7. Return the filled file and run the pipeline with STAGES=kappa SHEET=<path> to compute agreement.",
    ]:
        wi.append([line])
    wi.column_dimensions["A"].width = 120
    path = os.path.join(OUT, "screening_%s.xlsx" % TODAY)
    wb.save(path)
    log("workbook saved: %s (%d records, %d for double screening)" % (path, len(df), k))
    # search log template for the manual database searches
    tmpl = pd.DataFrame([dict(source=s, date="", query="(paste the exact string)", field="", years="%d-%d" % (YEAR_FROM, YEAR_TO),
                              filters="", hits="", exported="") for s in
                         ("Scopus", "Web of Science", "IEEE Xplore", "ACM DL")])
    api_log = os.path.join(OUT, "search_log_api.csv")
    if os.path.exists(api_log) and os.path.getsize(api_log) > 5:
        tmpl = pd.concat([tmpl, pd.read_csv(api_log)], ignore_index=True)
    tmpl.to_csv(os.path.join(OUT, "search_log.csv"), index=False)


# ----------------------------------------------------------------------------- agreement
def cohen_kappa(a, b):
    labs = sorted(set(a) | set(b))
    n = len(a)
    if n == 0:
        return float("nan")
    po = np.mean([x == y for x, y in zip(a, b)])
    pe = sum((np.mean([x == l for x in a])) * (np.mean([y == l for y in b])) for l in labs)
    return float((po - pe) / (1 - pe)) if pe < 1 else 1.0


def stage_kappa():
    sp = os.environ.get("SHEET") or sorted(glob.glob(os.path.join(DATA_DIR, "**", "screening*.xlsx"), recursive=True) +
                                           glob.glob(os.path.join(OUT, "screening*.xlsx")))[-1]
    df = pd.read_excel(sp, sheet_name="screening", dtype=str).fillna("")
    norm = lambda s: s.strip().lower()
    both = df[(df.S1_decision.map(norm) != "") & (df.S2_decision.map(norm) != "")]
    a, b = list(both.S1_decision.map(norm)), list(both.S2_decision.map(norm))
    k3 = cohen_kappa(a, b)
    bin_ = lambda x: "exclude" if x == "exclude" else "keep"
    k2 = cohen_kappa([bin_(x) for x in a], [bin_(x) for x in b])
    res = dict(sheet=os.path.basename(sp), n_double=len(both), kappa_3class=round(k3, 3), kappa_keep_vs_exclude=round(k2, 3),
               percent_agreement=round(float(np.mean([x == y for x, y in zip(a, b)])), 3))
    conf = both[[x != y for x, y in zip(a, b)]][["rec_id", "title", "S1_decision", "S2_decision"]]
    conf.to_csv(os.path.join(OUT, "conflicts.csv"), index=False)
    final = df.final_decision.map(norm).where(df.final_decision.map(norm) != "", df.S1_decision.map(norm))
    res["final_counts"] = final.value_counts().to_dict()
    if "ai_p_include" in df.columns and (df.ai_p_include != "").any():
        keep = final.isin(["include", "maybe"])
        p = pd.to_numeric(df.ai_p_include, errors="coerce")
        for thr in (0.15, 0.5):
            pred = p >= thr
            res["ai_recall_at_%.2f" % thr] = round(float((pred & keep).sum() / max(1, keep.sum())), 3)
            res["ai_workload_saved_at_%.2f" % thr] = round(float(1 - pred.mean()), 3)
        res["ai_kappa_vs_human"] = round(cohen_kappa(list(np.where(p >= 0.5, "keep", "exclude")),
                                                     [bin_(x) for x in final]), 3)
    json.dump(res, open(os.path.join(OUT, "agreement.json"), "w"), indent=1)
    log("agreement:", json.dumps(res))


# ----------------------------------------------------------------------------- main
def package():
    zp = "/kaggle/working/survey_screening_%s.zip" % TODAY if ON_KAGGLE else os.path.join(OUT, "..", "survey_screening_%s.zip" % TODAY)
    with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
        for f in glob.glob(os.path.join(OUT, "*")):
            if os.path.isfile(f):
                z.write(f, os.path.basename(f))
        z.write(os.path.abspath(__file__), "code_snapshot.py")
    log("results: %s" % zp)


if __name__ == "__main__":
    log("=" * 60)
    log("survey pipeline | stages %s | data %s | FAST %s" % (STAGES, DATA_DIR, FAST))
    for st, fn in (("harvest", stage_harvest), ("merge", stage_merge), ("llm", stage_llm), ("sheet", stage_sheet),
                   ("kappa", stage_kappa)):
        if st in STAGES:
            try:
                fn()
            except SystemExit:
                raise
            except Exception as e:
                log("!! stage %s failed: %s" % (st, e)); traceback.print_exc()
                if st in ("merge",):
                    raise
    package()
    log("finished")
